"""Check the actual proposal worker protocol in staged and installed packages.

No imports from scandocument: passing --worker always exercises the executable.
Synthetic, bounded fixtures cover DOCX, PDF, preview, ZIP and rejected private DTOs.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import os
import signal
import subprocess
import tempfile
import time
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from PIL import Image
from PIL.PngImagePlugin import PngInfo
from pypdf import PdfReader

REPOSITORY = Path(__file__).resolve().parents[1]
SENTINEL = "SBK_PRIVATE_SENTINEL_MUST_NOT_EXPORT"
ATTACHMENT = "PUBLIC_ATTACHMENT_ACCEPTED\n".encode()


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def document_fixture() -> dict:
    party = {"name": "", "shortName": "", "inn": "", "kpp": "", "ogrn": "",
             "address": "Москва", "contact": "", "paymentDetails": ""}
    contact = {"fullName": "", "position": "", "phone": "", "email": ""}
    taxable = {"title": "Настройка системы", "description": "Передача документации", "unit": "усл.",
               "quantity": "1", "unitPrice": "12000", "priceBasis": "gross", "discountPercent": "0",
               "tax": {"kind": "vat", "rate": 20}, "netMinor": "1000000", "vatMinor": "200000", "grossMinor": "1200000"}
    exempt = {"title": "Обучение пользователей", "description": "Практическое занятие", "unit": "ч",
              "quantity": "2.5", "unitPrice": "100", "priceBasis": "net", "discountPercent": "10",
              "tax": {"kind": "none"}, "netMinor": "22500", "vatMinor": "0", "grossMinor": "22500"}
    return {
        "schemaVersion": 1, "number": "КП-20300101-QA", "revision": 1,
        "documentDate": "2030-01-01", "validUntil": "2030-01-31", "currency": "RUB",
        "title": "Внедрение информационной системы",
        "issuer": {**party, "name": "Тестовый исполнитель"},
        "recipient": {**party, "name": "Тестовый заказчик"}, "addressee": contact.copy(), "contact": contact.copy(),
        "lines": [taxable, exempt],
        "totals": {"pricingVersion": "proposal-pricing/1", "netMinor": "1022500", "vatMinor": "200000", "grossMinor": "1222500",
                   "byTax": [{key: line[key] for key in ("tax", "netMinor", "vatMinor", "grossMinor")} for line in (taxable, exempt)]},
        "terms": {"introduction": "Предлагаем работы и обучение в согласованном объёме.",
                  "delivery": "30 календарных дней", "payment": "После приёмки", "conclusion": "Готовы обсудить состав работ."},
        "layout": {"style": "compact", "accentColor": "#245C48",
                   "show": {"address": True, "requisites": True, "contact": True, "signer": True},
                   "footer": "Тестовое предложение"},
        "attachments": [],
    }


def inspect_docx(data: bytes, forbidden: list[str]) -> None:
    with ZipFile(io.BytesIO(data)) as archive:
        assert "word/document.xml" in archive.namelist(), "Result is not a DOCX"
        for name in archive.namelist():
            content = archive.read(name)
            for token in forbidden:
                assert token.encode() not in content, f"Private value leaked into DOCX {name}"
    document = Document(io.BytesIO(data))
    text = " ".join([paragraph.text for paragraph in document.paragraphs] +
                    [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    assert len(document.tables) == 1 and len(document.tables[0].rows) == 3
    assert len(document.inline_shapes) == 1, "Selected logo not rendered"
    assert abs(document.sections[0].page_width.mm - 210) < 0.1
    for expected in ("Тестовый исполнитель", "Тестовый заказчик", "Настройка системы", "Обучение пользователей", "12 225,00"):
        assert expected in text, f"Missing DOCX content: {expected}"
    assert not document.core_properties.author and not document.core_properties.last_modified_by


def inspect_pdf(data: bytes, forbidden: list[str]) -> int:
    assert data.startswith(b"%PDF-"), "Result is not a PDF"
    reader = PdfReader(io.BytesIO(data))
    assert 1 <= len(reader.pages) <= 3, "Small proposal has an unexpected page count"
    text = " ".join(" ".join((page.extract_text() or "").split()) for page in reader.pages)
    for expected in ("Коммерческое предложение", "Настройка системы", "Обучение пользователей", "12 225,00"):
        assert expected in text, f"Missing PDF content: {expected}"
    searchable = text + str(reader.metadata) + str(reader.xmp_metadata)
    for token in forbidden:
        assert token not in searchable and token.encode() not in data, "Private value leaked into PDF"
    return len(reader.pages)


def run_smoke(command: list[str], runtime: Path, directory: Path, *, office_override: Path | None = None) -> dict:
    """office_override is for explicit source-protocol QA, never exposed by the packaged CLI."""
    directory = directory.resolve()
    if os.name == "nt":
        # Match Rust Path::canonicalize in the native proposal bridge. A plain
        # tempfile path hides Office URI bugs involving Windows verbatim paths.
        windows_path = str(directory)
        if not windows_path.startswith("\\\\?\\"):
            windows_path = ("\\\\?\\UNC\\" + windows_path[2:]
                            if windows_path.startswith("\\\\") else "\\\\?\\" + windows_path)
        directory = Path(windows_path)
    environment = dict(os.environ)
    for name in ("SCANDOCUMENT_SOFFICE", "PYTHONPATH", "PYTHONHOME"):
        environment.pop(name, None)
    environment["SCANDOCUMENT_RESOURCE_ROOT"] = str(runtime.resolve())
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    if office_override is not None:
        environment["SCANDOCUMENT_SOFFICE"] = str(office_override)
    checks = []
    deadline = time.monotonic() + 600
    forbidden = [SENTINEL, str(directory), directory.as_posix()]

    def invoke(label: str, config: dict, success: bool = True) -> dict:
        workdir = Path(config["workdir"])
        config_path = directory / f"{label}.json"
        config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        started = time.monotonic()
        remaining = deadline - started
        assert remaining > 0, "Proposal smoke exceeded its 10 minute budget"
        process = subprocess.Popen([*command, "proposal", "--config", str(config_path)], env=environment,
                                   stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8",
                                   start_new_session=os.name != "nt")
        try:
            stdout, stderr = process.communicate(timeout=min(180, remaining))
        except subprocess.TimeoutExpired:
            (workdir / "cancel.requested").write_text("cancel", encoding="utf-8")
            try:
                stdout, stderr = process.communicate(timeout=5)
            except subprocess.TimeoutExpired:
                if os.name == "nt":
                    subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True, timeout=15, check=False)
                else:
                    os.killpg(process.pid, signal.SIGKILL)
                process.kill()
                stdout, stderr = process.communicate()
            (directory / f"{label}.jsonl").write_text(stdout, encoding="utf-8")
            (directory / f"{label}.stderr.log").write_text(stderr, encoding="utf-8")
            raise AssertionError(f"{label}: proposal worker timed out") from None
        (directory / f"{label}.jsonl").write_text(stdout, encoding="utf-8")
        (directory / f"{label}.stderr.log").write_text(stderr, encoding="utf-8")
        events = [json.loads(line) for line in stdout.splitlines() if line.strip()]
        assert events, f"{label}: no protocol output; {stderr[-1500:]}"
        event = events[-1]
        if success:
            assert process.returncode == 0 and event.get("type") == "complete", f"{label}: {event}"
            extension = "pdf" if config["format"] == "preview" else config["format"]
            output = workdir / f"proposal.{extension}"
            assert Path(event["outputPath"]).resolve() == output and output.is_file()
            data = output.read_bytes()
            assert len(data) == event["outputBytes"] and digest(data) == event["sha256"]
        else:
            assert process.returncode != 0 and event.get("type") == "error", f"{label}: expected rejection"
            assert not list(workdir.glob("proposal.*")), f"{label}: rejected input produced an output"
        checks.append({"check": label, "exitCode": process.returncode, "pageCount": event.get("pageCount"),
                       "elapsedSeconds": round(time.monotonic() - started, 3)})
        return event

    for output_format in ("docx", "pdf", "preview", "zip"):
        workdir = directory / output_format
        assets = workdir / "assets"
        assets.mkdir(parents=True)
        attachment = assets / digest(ATTACHMENT)
        attachment.write_bytes(ATTACHMENT)
        logo = assets / "input-logo.png"
        metadata = PngInfo()
        metadata.add_text("internal", SENTINEL)
        with Image.new("RGB", (240, 60), "#245C48") as image:
            image.save(logo, pnginfo=metadata)
        logo_bytes = logo.read_bytes()
        # An unrelated private file must not enter the ZIP merely by sharing the folder.
        (assets / "unselected-private.txt").write_text(SENTINEL, encoding="utf-8")
        document = document_fixture()
        selected = {"fileName": "Условия.txt", "sizeBytes": len(ATTACHMENT), "sha256": digest(ATTACHMENT), "mimeType": "text/plain"}
        document["attachments"] = [selected]
        document["layout"]["logo"] = {"fileName": "Логотип.png", "sizeBytes": len(logo_bytes), "sha256": digest(logo_bytes), "mimeType": "image/png"}
        config = {"workdir": str(workdir), "format": output_format, "document": document,
                  "assets": [{"path": str(attachment), "sha256": digest(ATTACHMENT)}, {"path": str(logo), "sha256": digest(logo_bytes)}]}
        event = invoke(output_format, config)
        inspect_docx((workdir / "proposal.docx").read_bytes(), forbidden)
        if output_format != "docx":
            pages = inspect_pdf((workdir / "proposal.pdf").read_bytes(), forbidden)
            assert event["pageCount"] == pages
        if output_format == "preview":
            assert len(event["previewPages"]) == pages
            for index, path in enumerate(event["previewPages"], 1):
                image_path = workdir / f"page-{index}.png"
                assert Path(path).resolve() == image_path
                with Image.open(image_path) as image:
                    assert image.format == "PNG" and min(image.size) >= 500
                    image.load()
                    assert any(low != high for low, high in image.convert("RGB").getextrema()), "Preview is blank"
        if output_format == "zip":
            with ZipFile(workdir / "proposal.zip") as archive:
                expected = {"Коммерческое предложение.docx", "Коммерческое предложение.pdf", "Приложения/001-Условия.txt", "Перечень.json"}
                assert set(archive.namelist()) == expected and len(archive.infolist()) == len(expected)
                assert archive.read("Приложения/001-Условия.txt") == ATTACHMENT
                inspect_docx(archive.read("Коммерческое предложение.docx"), forbidden)
                inspect_pdf(archive.read("Коммерческое предложение.pdf"), forbidden)
                manifest = json.loads(archive.read("Перечень.json"))
                assert manifest["files"] == [{"name": "Приложения/001-Условия.txt", "sizeBytes": len(ATTACHMENT), "sha256": digest(ATTACHMENT)}]
                for token in forbidden:
                    assert token not in json.dumps(manifest), "Private value leaked into ZIP manifest"
        assert attachment.read_bytes() == ATTACHMENT and logo.read_bytes() == logo_bytes

    for label, location in (("private-root", "root"), ("private-line", "line")):
        workdir = directory / label
        workdir.mkdir()
        document = copy.deepcopy(document_fixture())
        target = document if location == "root" else document["lines"][0]
        target["internalNote"] = SENTINEL
        invoke(label, {"workdir": str(workdir), "format": "docx", "document": document, "assets": []}, success=False)
    return {"status": "passed", "command": command, "runtime": str(runtime.resolve()), "checks": checks,
            "scope": "Compiled executable protocol" if office_override is None else "Source protocol with explicit QA Office override; not packaged evidence"}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", type=Path, required=True)
    parser.add_argument("--runtime", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    assert args.worker.is_file() and args.runtime.is_dir(), "Worker or runtime missing"
    args.output_dir.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.output_dir.resolve()))
    try:
        report = run_smoke([str(args.worker.resolve())], args.runtime, directory)
    except Exception as error:
        (directory / "report.json").write_text(json.dumps({"status": "failed", "error": str(error)}, ensure_ascii=False, indent=2), encoding="utf-8")
        raise
    (directory / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
