"""Measure end-to-end scanner CLI latency using generated, non-user documents.

Cold means an empty application cache, NOT flushed OS/disk/LibreOffice caches.
Every operation starts a new worker, matching the current application lifecycle.
No source documents, existing caches or databases are modified or deleted.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import platform
import statistics
import subprocess
import sys
import tempfile
import time
from pathlib import Path


REPOSITORY = Path(__file__).resolve().parents[1]


def positive_integer(value: str) -> int:
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("Must be a positive integer")
    return number


def summarize(samples: list[dict]) -> dict:
    groups: dict[str, list[float]] = {}
    for sample in samples:
        groups.setdefault(sample["scenario"], []).append(sample["elapsedSeconds"])
    return {
        name: {"samples": len(values), "medianSeconds": round(statistics.median(values), 4),
               "p95Seconds": round(sorted(values)[math.ceil(len(values) * 0.95) - 1], 4),
               "minSeconds": round(min(values), 4), "maxSeconds": round(max(values), 4)}
        for name, values in groups.items()
    }


def source_identity(directory: Path) -> str:
    digest = hashlib.sha256()
    for path in sorted((directory / "scandocument").rglob("*.py")):
        digest.update(path.relative_to(directory).as_posix().encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def read_terminal_event(stdout: str, operation: str) -> dict:
    events = [json.loads(line) for line in stdout.splitlines() if line.strip()]
    expected = {"preview": "preview", "preparePreview": "prepared", "process": "complete"}[operation]
    if not events or not isinstance(events[-1], dict) or events[-1].get("type") != expected:
        raise ValueError(f"Expected terminal {expected} event; received {events[-1:]!r}")
    if events[-1].get("protocolVersion") != 2:
        raise ValueError("Unexpected worker protocol version")
    return events[-1]


def make_fixtures(directory: Path, pages: int, formats: list[str]) -> list[Path]:
    """Generate deterministic text/geometry, without opening any user document."""
    paths = []
    if "pdf" in formats:
        from reportlab.pdfgen.canvas import Canvas

        path = directory / "synthetic.pdf"
        canvas = Canvas(str(path), pagesize=(595.28, 841.89), invariant=True)
        for index in range(pages):
            canvas.setFont("Helvetica", 18)
            canvas.drawString(45, 785, f"SBK scanner benchmark / page {index + 1}")
            canvas.setFont("Helvetica", 11)
            for row in range(36):
                canvas.drawString(45, 745 - row * 17, f"Synthetic public test content {index + 1:03d}:{row + 1:02d}")
            canvas.setFillColorRGB((index % 3) / 3, 0.4, 0.7)
            canvas.rect(45, 65, 120, 40, fill=1)
            canvas.showPage()
        canvas.save()
        paths.append(path)
    if "docx" in formats:
        from docx import Document

        path = directory / "synthetic.docx"
        document = Document()
        for index in range(pages):
            if index:
                document.add_page_break()
            document.add_heading(f"SBK scanner benchmark / page {index + 1}", level=1)
            document.add_paragraph("Synthetic public test content. No user records or attachments.")
            table = document.add_table(rows=8, cols=3)
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    cell.text = f"Page {index + 1}, row {row_index + 1}, column {column_index + 1}"
        document.save(path)
        paths.append(path)
    return paths


def case_plan(source: Path, cache: Path, outputs: Path, pages: int, dpi: int) -> list[tuple[str, str, dict]]:
    base = {"protocolVersion": 2, "inputPath": str(source), "preset": "Офисный скан", "seed": 42,
            "previewCacheDir": str(cache), "settings": {"dpi": dpi, "jpeg_quality": 50},
            "ocrEnabled": False, "outputPolicy": "no-clobber"}
    cases: list[tuple[str, str, dict]] = []

    def preview(name: str, page: int) -> None:
        cases.append((name, "preview", {**base, "pageIndex": page, "outputPath": str(outputs / f"{name}.png")}))

    preview("first-page-cold", 0)
    preview("first-page-warm", 0)
    if pages > 1:
        preview("next-page-cold-raster", 1)
        preview("next-page-warm-raster", 1)
    if pages > 2:
        indices = list(range(2, min(pages, 5)))
        cases.append(("prepare-neighbors", "preparePreview", {
            **base, "pageIndices": indices,
            "outputPaths": [str(outputs / f"prepared-{index}.png") for index in indices],
        }))
        preview("prepared-neighbor", 2)
    cases.append(("export-after-preview", "process", {**base, "outputPath": str(outputs / "prepared-export.pdf")}))
    if source.suffix.lower() == ".docx":
        cases.append(("export-empty-document-cache", "process", {
            **base, "previewCacheDir": str(cache.with_name(cache.name + "-export-cold")),
            "outputPath": str(outputs / "cold-export.pdf"),
        }))
    return cases


def verify_result(event: dict, operation: str, config: dict, pages: int) -> None:
    """Keep validation outside the timed region; reject misleading fast failures."""
    if operation == "process":
        from pypdf import PdfReader

        path = Path(config["outputPath"])
        if Path(event["outputPath"]) != path or not path.is_file() or len(PdfReader(path).pages) != pages:
            raise ValueError("Export output or page count does not match the generated fixture")
        return
    from PIL import Image

    previews = event["previews"] if operation == "preparePreview" else [event]
    indices = config["pageIndices"] if operation == "preparePreview" else [config["pageIndex"]]
    paths = config["outputPaths"] if operation == "preparePreview" else [config["outputPath"]]
    if len(previews) != len(paths):
        raise ValueError("Incorrect number of prepared previews")
    for preview, index, output in zip(previews, indices, paths, strict=True):
        if preview["pageCount"] != pages or preview["pageIndex"] != index or preview["outputPath"] != output:
            raise ValueError("Preview page identity does not match the generated fixture")
        output_path = Path(output)
        expected_original = output_path.with_name(output_path.stem + ".original.png")
        if Path(preview["originalPath"]) != expected_original:
            raise ValueError("Original preview path does not match the generated fixture")
        for path in [output_path, expected_original]:
            with Image.open(path) as image:
                if image.format != "PNG" or min(image.size) < 1:
                    raise ValueError("Invalid preview image")
                image.verify()


def run_case(command: list[str], environment: dict, directory: Path, name: str,
             operation: str, config: dict, pages: int, timeout: int) -> dict:
    config_path = directory / f"{name}.json"
    config_path.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")
    stdout_path = directory / f"{name}.jsonl"
    stderr_path = directory / f"{name}.stderr.log"
    started = time.perf_counter()
    with stdout_path.open("w", encoding="utf-8") as stdout, stderr_path.open("w", encoding="utf-8") as stderr:
        result = subprocess.run([*command, operation, "--config", str(config_path)], env=environment,
                                stdout=stdout, stderr=stderr, timeout=timeout, check=False)
    elapsed = time.perf_counter() - started
    if result.returncode:
        raise RuntimeError(f"{name}: worker exit {result.returncode}; see {stdout_path} and {stderr_path}")
    event = read_terminal_event(stdout_path.read_text(encoding="utf-8"), operation)
    verify_result(event, operation, config, pages)
    return {"scenario": f"{Path(config['inputPath']).suffix[1:]}:{name}", "operation": operation,
            "elapsedSeconds": round(elapsed, 4), "exitCode": result.returncode, "sourceFingerprint": event.get("sourceFingerprint"),
            "outputBytes": event.get("outputBytes"), "configPath": str(config_path)}


def create_run_directory(parent: Path, prefix: str) -> Path:
    parent.mkdir(parents=True, exist_ok=True)
    return Path(tempfile.mkdtemp(prefix=prefix, dir=parent.resolve()))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--worker", type=Path, help="Packaged worker binary; no worker-source imports")
    mode.add_argument("--python", default=sys.executable, help="Python with scanner dependencies for source-worker mode")
    parser.add_argument("--source-dir", type=Path, help="Alternate directory containing scandocument/, for baseline comparisons")
    parser.add_argument("--runtime", type=Path, help="Runtime directory containing resources/ (required for packaged DOCX)")
    parser.add_argument("--output-dir", type=Path, default=REPOSITORY / "scanner-worker/tests/output/benchmark")
    parser.add_argument("--input-root", type=Path, help="Optional shared-folder root for new synthetic input files")
    parser.add_argument("--cache-root", type=Path, help="Optional cache root, to compare local and shared cache I/O")
    parser.add_argument("--pages", type=positive_integer, default=3)
    parser.add_argument("--repeats", type=positive_integer, default=3)
    parser.add_argument("--timeout", type=positive_integer, default=300, help="Per-operation timeout in seconds")
    parser.add_argument("--dpi", type=int, choices=[96, 120, 150, 200, 300], default=120)
    parser.add_argument("--formats", nargs="+", choices=["pdf", "docx"], default=["pdf", "docx"])
    parser.add_argument("--label", default="", help="Human-readable build/environment label")
    args = parser.parse_args()
    if args.pages > 1000:
        parser.error("Use no more than 1000 synthetic pages")
    if args.worker and not args.worker.is_file():
        parser.error(f"Worker not found: {args.worker}")
    if args.worker and args.source_dir:
        parser.error("--source-dir cannot be combined with --worker")
    source_directory = (args.source_dir or REPOSITORY / "scanner-worker/src").resolve()
    if not args.worker and not (source_directory / "scandocument/worker_cli.py").is_file():
        parser.error("--source-dir must contain scandocument/worker_cli.py")
    if args.runtime and not (args.runtime / "resources").is_dir():
        parser.error("--runtime must contain a resources directory")
    environment = dict(os.environ)
    if args.runtime:
        environment["SCANDOCUMENT_RESOURCE_ROOT"] = str(args.runtime.resolve())
    if args.worker:
        command = [str(args.worker.resolve())]
        # A compiled benchmark must not silently import from a developer checkout.
        environment.pop("PYTHONPATH", None)
    else:
        command = [args.python, "-m", "scandocument.worker_cli"]
        environment["PYTHONPATH"] = str(source_directory)
    directory = create_run_directory(args.output_dir, "run-")
    inputs = create_run_directory(args.input_root or directory, "synthetic-inputs-")
    caches = create_run_directory(args.cache_root or directory, "isolated-caches-")
    report = {
        "status": "running", "label": args.label, "mode": "packaged" if args.worker else "source",
        "command": command, "host": platform.platform(), "python": sys.version.split()[0],
        "pages": args.pages, "repeats": args.repeats, "dpi": args.dpi,
        "inputDirectory": str(inputs), "cacheDirectory": str(caches), "outputDirectory": str(directory),
        "runtime": environment.get("SCANDOCUMENT_RESOURCE_ROOT"), "samples": [],
        "officeOverride": environment.get("SCANDOCUMENT_SOFFICE"),
        "sourceDirectory": str(source_directory) if not args.worker else None,
        "sourceSha256": source_identity(source_directory) if not args.worker else None,
        "workerSha256": hashlib.sha256(args.worker.read_bytes()).hexdigest() if args.worker else None,
        "scope": "End-to-end CLI, including worker startup/shutdown. Not UI latency or peak-memory measurement.",
        "coldDefinition": "Empty application cache; OS/filesystem/LibreOffice caches are not flushed.",
        "p95Definition": "Nearest-rank observed percentile; small sample counts are exploratory, not an SLO.",
        "retention": "Only new isolated directories were written. Nothing is automatically deleted.",
    }
    report_path = directory / "report.json"
    try:
        sources = make_fixtures(inputs, args.pages, list(dict.fromkeys(args.formats)))
        original_hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in sources}
        report["inputSha256"] = {str(path): digest for path, digest in original_hashes.items()}
        for repeat in range(args.repeats):
            for source in sources:
                kind = source.suffix[1:]
                outputs = directory / f"{kind}-{repeat + 1}"
                outputs.mkdir()
                for name, operation, config in case_plan(source, caches / f"{kind}-{repeat + 1}", outputs, args.pages, args.dpi):
                    sample = run_case(command, environment, outputs, name, operation, config, args.pages, args.timeout)
                    sample["repeat"] = repeat + 1
                    report["samples"].append(sample)
                    print(json.dumps(sample, ensure_ascii=False), flush=True)
        if any(hashlib.sha256(path.read_bytes()).hexdigest() != digest for path, digest in original_hashes.items()):
            raise ValueError("A generated source was unexpectedly modified")
        if not args.worker and source_identity(source_directory) != report["sourceSha256"]:
            raise ValueError("Worker source changed during the benchmark; rerun with a stable checkout")
        report["status"] = "passed"
    except Exception as error:
        report.update(status="failed", error=str(error))
        raise
    finally:
        report["summary"] = summarize(report["samples"])
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Benchmark report: {report_path}", flush=True)


if __name__ == "__main__":
    main()
