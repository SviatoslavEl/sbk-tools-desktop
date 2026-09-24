"""Exercise a real scanner executable using isolated PDF/DOCX fixtures.

The same script runs against the built sidecar, the final Mac bundle, and the
installed Windows copy. It intentionally never imports the worker's source.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import tempfile
import time
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter


REPOSITORY = Path(__file__).resolve().parents[1]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run_smoke(worker: Path, runtime: Path, directory: Path) -> dict:
    fixtures = REPOSITORY / "scanner-worker/tests/fixtures"
    originals = {path: sha256(path) for path in [fixtures / "simple.pdf", fixtures / "simple.docx"]}
    pdf = directory / "three-pages.pdf"
    writer = PdfWriter()
    reader = PdfReader(fixtures / "simple.pdf")
    for _ in range(3):
        writer.add_page(reader.pages[0])
    with pdf.open("wb") as stream:
        writer.write(stream)
    docx = directory / "three-pages.docx"
    document = Document(fixtures / "simple.docx")
    for number in [2, 3]:
        document.add_page_break()
        document.add_heading(f"SCANNER PREPARATION QA PAGE {number}", level=1)
        document.add_paragraph("Synthetic offline preview test. No user data.")
    document.save(docx)
    environment = dict(os.environ)
    environment["SCANDOCUMENT_RESOURCE_ROOT"] = str(runtime.resolve())
    reports = []

    def run(operation: str, config: dict, label: str, expect_success: bool = True) -> tuple[dict, list[dict]]:
        config_path = directory / f"{label}.json"
        config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        started = time.perf_counter()
        result = subprocess.run([str(worker.resolve()), operation, "--config", str(config_path)],
                                env=environment, capture_output=True, text=True, encoding="utf-8", timeout=300)
        elapsed = time.perf_counter() - started
        (directory / f"{label}.jsonl").write_text(result.stdout, encoding="utf-8")
        (directory / f"{label}.stderr.log").write_text(result.stderr, encoding="utf-8")
        events = [json.loads(line) for line in result.stdout.splitlines() if line.strip()]
        assert events, f"{label}: worker returned no events: {result.stderr[-1500:]}"
        if expect_success:
            assert result.returncode == 0, f"{label}: {events[-1]}"
        else:
            assert result.returncode != 0 and events[-1].get("type") == "error", f"{label}: expected rejection"
        reports.append({"check": label, "operation": operation, "exitCode": result.returncode,
                        "elapsedSeconds": round(elapsed, 4)})
        return events[-1], events

    def verify_preview(page: dict, expected_path: Path, expected_index: int, count: int, fingerprint: str) -> list[Path]:
        assert page["type"] == "preview" and page["protocolVersion"] == 2
        assert page["pageCount"] == count and page["pageIndex"] == expected_index
        assert page["sourceFingerprint"] == fingerprint
        expected = [expected_path, expected_path.with_name(f"{expected_path.stem}.original.png")]
        actual = [Path(page["outputPath"]), Path(page["originalPath"])]
        assert actual == expected, f"Unexpected output paths: {actual}"
        for path in actual:
            assert path.is_file() and path.stat().st_size > 8
            with Image.open(path) as image:
                assert image.format == "PNG" and all(0 < dimension <= 1400 for dimension in image.size)
                image.verify()
        return actual

    for source in [pdf, docx]:
        kind = source.suffix[1:]
        cache = directory / f"cache-{kind}"
        base = {"protocolVersion": 2, "inputPath": str(source), "preset": "Офисный скан", "seed": 42,
                "previewCacheDir": str(cache), "settings": {"dpi": 96, "jpeg_quality": 40}}
        first_path = directory / f"{kind}-first.png"
        first, _ = run("preview", {**base, "pageIndex": 0, "outputPath": str(first_path)}, f"{kind}-first")
        count = first["pageCount"]
        assert count >= 3, f"{kind}: fixture has too few pages"
        fingerprint = first["sourceFingerprint"]
        assert re.fullmatch(r"[0-9a-f]{64}", fingerprint)
        owned = verify_preview(first, first_path, 0, count, fingerprint)
        indices = [count - 1, 0, 1]
        paths = [directory / f"{kind}-prepared-{index}.png" for index in indices]
        config = {**base, "pageIndices": indices, "outputPaths": [str(path) for path in paths]}
        prepared, events = run("preparePreview", config, f"{kind}-prepared")
        assert prepared["type"] == "prepared" and prepared["protocolVersion"] == 2
        assert prepared["pageCount"] == count and prepared["sourceFingerprint"] == fingerprint
        assert len(prepared["previews"]) == 3
        progress = [event for event in events if event["type"] == "progress"]
        assert [event["currentPage"] for event in progress] == [1, 2, 3]
        assert [event["percent"] for event in progress] == [33, 67, 100]
        assert all(event["totalPages"] == 3 for event in progress)
        for page, path, index in zip(prepared["previews"], paths, indices, strict=True):
            owned.extend(verify_preview(page, path, index, count, fingerprint))
        assert len({page["estimatedOutputBytes"] for page in prepared["previews"]}) == 1
        expected_pixels = {path.name: sha256(path) for path in owned}
        # These are the same disposable job-owned copies the UI removes after
        # reading. Persistent cache entries must survive independently.
        for path in owned:
            assert path.parent == directory
            path.unlink()
        assert all(not path.exists() for path in owned)
        repeated, _ = run("preparePreview", config, f"{kind}-repeated-after-cleanup")
        for page, path, index in zip(repeated["previews"], paths, indices, strict=True):
            for copy in verify_preview(page, path, index, count, fingerprint):
                assert sha256(copy) == expected_pixels[copy.name], "Cached output pixels changed"
                copy.unlink()
        invalid_paths = [directory / f"{kind}-invalid-{index}.png" for index in range(2)]
        run("preparePreview", {**base, "pageIndices": [0, 0], "outputPaths": [str(p) for p in invalid_paths]},
            f"{kind}-invalid-batch", expect_success=False)
        assert not any(path.exists() for path in invalid_paths)
        cached_pdf = None
        cache_digest = None
        cache_warning = "SCANNER PREPARATION QA: prepared DOCX conversion reused"
        if kind == "docx":
            conversions = list(cache.glob(f"docx-*-{fingerprint}.pdf"))
            assert len(conversions) == 1, "Expected one prepared DOCX conversion"
            cached_pdf = conversions[0]
            cache_digest = sha256(cached_pdf)
            # Modify only metadata of our synthetic, per-run cache. The marker
            # proves export uses this conversion; a fresh conversion would lose
            # it. Never alter the installed runtime or repository fixtures.
            warning_path = cached_pdf.with_suffix(".warnings.json")
            warnings = json.loads(warning_path.read_text(encoding="utf-8"))
            assert isinstance(warnings, list)
            warning_path.write_text(json.dumps([*warnings, cache_warning]), encoding="utf-8")
        output = directory / f"{kind}-unchanged.pdf"
        complete, _ = run("process", {**base, "outputPath": str(output), "expectedSourceFingerprint": fingerprint},
                          f"{kind}-unchanged-process")
        assert complete["type"] == "complete" and len(PdfReader(output).pages) == count
        if cached_pdf is not None:
            assert cache_warning in complete["warnings"], "Export did not reuse the prepared DOCX conversion"
            assert sha256(cached_pdf) == cache_digest, "Export replaced the prepared DOCX conversion"
            merged_path = directory / "docx-cached-merge.pdf"
            merged, _ = run("merge", {**base, "inputPaths": [str(source), str(pdf)],
                                     "outputPath": str(merged_path)}, "docx-cached-merge")
            assert merged["type"] == "complete"
            assert len(PdfReader(merged_path).pages) == count + len(PdfReader(pdf).pages)
            assert cache_warning in merged["warnings"], "Merge did not reuse the prepared DOCX conversion"
            assert sha256(cached_pdf) == cache_digest, "Merge replaced the prepared DOCX conversion"
        source_metadata = source.stat()
        os.utime(source, ns=(source_metadata.st_atime_ns, source_metadata.st_mtime_ns + 5_000_000_000))
        rejected_path = directory / f"{kind}-stale-must-not-exist.pdf"
        rejected, _ = run("process", {**base, "outputPath": str(rejected_path), "expectedSourceFingerprint": fingerprint},
                          f"{kind}-changed-source-rejected", expect_success=False)
        assert "изменился" in rejected.get("message", "") and not rejected_path.exists()
        assert len(list(cache.glob("raster-*.png"))) <= 64
        assert sum(path.stat().st_size for path in cache.glob("raster-*.png")) <= 128 * 1024 * 1024
        assert not list(directory.glob(f"{kind}-*.png")), "Job-owned preview copies were not cleaned"

    assert all(sha256(path) == digest for path, digest in originals.items()), "Repository fixture changed"
    return {"status": "passed", "worker": str(worker.resolve()), "runtime": str(runtime.resolve()),
            "outputDirectory": str(directory), "checks": reports,
            "scope": "Real scanner executable CLI; not a Windows GUI test or Rust cancellation test."}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", type=Path, required=True)
    parser.add_argument("--runtime", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=REPOSITORY / "scanner-worker/tests/output/preview-preparation")
    args = parser.parse_args()
    assert args.worker.is_file(), f"Worker missing: {args.worker}"
    assert args.runtime.is_dir(), f"Runtime missing: {args.runtime}"
    args.output_dir.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="run-", dir=args.output_dir.resolve()))
    try:
        report = run_smoke(args.worker, args.runtime, directory)
    except Exception as error:
        (directory / "report.json").write_text(json.dumps({"status": "failed", "error": str(error)}, indent=2), encoding="utf-8")
        raise
    (directory / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
