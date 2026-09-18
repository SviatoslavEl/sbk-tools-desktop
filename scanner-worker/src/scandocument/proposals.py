"""Customer-only commercial proposals. This renderer never receives record payloads."""

from __future__ import annotations

import hashlib
import json
import re
import signal
import threading
from datetime import date
from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageOps

from scandocument.resources import find_soffice
from scandocument.errors import CancelledError

MAX_MINOR = 9007199254740991
RATES = (0, 5, 7, 10, 11, 20, 22)


def _object(value, required: str, optional: str = "") -> dict:
    needed, allowed = set(required.split()), set((required + " " + optional).split())
    if not isinstance(value, dict) or not needed <= value.keys() or not value.keys() <= allowed:
        raise ValueError("Некорректная публичная схема КП; внутренние поля запрещены")
    return value


def _text(value, maximum=20000, required=False) -> str:
    if not isinstance(value, str) or len(value) > maximum or (required and not value.strip()):
        raise ValueError("КП содержит пустой или слишком длинный обязательный текст")
    if any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise ValueError("Текст КП содержит недопустимые управляющие символы")
    return value


def _date(value: str, optional=False):
    if optional and value == "":
        return
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ValueError("Некорректная дата КП")
    date.fromisoformat(value)


def _decimal(value: str, precision: int, maximum: str) -> Decimal:
    if not isinstance(value, str) or len(value) > 40 or not re.fullmatch(r"\d+(?:\.\d+)?", value):
        raise ValueError("Некорректное десятичное число в КП")
    if len(value.partition(".")[2]) > precision:
        raise ValueError("Превышена точность числа в КП")
    result = Decimal(value)
    if not 0 <= result <= Decimal(maximum):
        raise ValueError("Превышен диапазон числа в КП")
    return result


def _tax(value):
    _object(value, "kind", "rate")
    if value == {"kind": "none"}:
        return
    if value.get("kind") != "vat" or type(value.get("rate")) is not int or value["rate"] not in RATES:
        raise ValueError("Некорректная ставка НДС")


def _minor(value):
    if not isinstance(value, str) or not re.fullmatch(r"0|[1-9]\d{0,15}", value) or int(value) > MAX_MINOR:
        raise ValueError("Превышен диапазон денежной суммы КП")


def _asset(value):
    _object(value, "fileName sizeBytes sha256 mimeType")
    _text(value["fileName"], 255, True)
    if re.search(r'[\\/<>:"|?*\x00-\x1f]', value["fileName"]) or value["fileName"] in (".", ".."):
        raise ValueError("Некорректное имя приложения КП")
    _text(value["mimeType"], 100)
    if type(value["sizeBytes"]) is not int or not 0 <= value["sizeBytes"] <= 512 * 1024 * 1024:
        raise ValueError("Приложение КП превышает 512 МБ")
    if not isinstance(value["sha256"], str) or not re.fullmatch("[0-9a-f]{64}", value["sha256"]):
        raise ValueError("Некорректный хеш приложения КП")


def validate_document(document: dict) -> dict:
    _object(
        document,
        "schemaVersion number revision documentDate title validUntil currency issuer recipient addressee contact lines totals terms layout attachments",
        "signer",
    )
    if (
        document["schemaVersion"] != 1
        or document["currency"] != "RUB"
        or type(document["revision"]) is not int
        or not 1 <= document["revision"] <= 100000
    ):
        raise ValueError("Неподдерживаемая версия или валюта КП")
    for key in ("number", "title"):
        _text(document[key], 500, True)
    for key in ("documentDate", "validUntil"):
        _date(document[key])
    if document["validUntil"] < document["documentDate"]:
        raise ValueError("Срок действия не может быть раньше даты КП")
    for key in ("issuer", "recipient"):
        party = _object(document[key], "name shortName inn kpp ogrn address contact paymentDetails")
        for name, value in party.items():
            _text(value, 500 if name in ("name", "shortName") else 20000, name == "name")
    for key in ("addressee", "contact", "signer"):
        if key not in document:
            continue
        contact = _object(
            document[key], "fullName position phone email" + (" basis issuedAt expiresAt" if key == "signer" else "")
        )
        for value in contact.values():
            _text(value)
        if key == "signer":
            _date(contact["issuedAt"], True)
            _date(contact["expiresAt"], True)
    lines = document["lines"]
    if not isinstance(lines, list) or not 1 <= len(lines) <= 1000:
        raise ValueError("В КП допустимо от 1 до 1000 позиций")
    totals = {"netMinor": 0, "vatMinor": 0, "grossMinor": 0}
    groups = {}
    with localcontext() as context:
        context.prec = 64
        for line in lines:
            _object(
                line,
                "title description unit quantity unitPrice priceBasis discountPercent tax netMinor vatMinor grossMinor",
            )
            _text(line["title"], 500, True)
            _text(line["description"])
            _text(line["unit"], 50, True)
            quantity = _decimal(line["quantity"], 6, "1000000000")
            price = _decimal(line["unitPrice"], 4, "1000000000000")
            discount = _decimal(line["discountPercent"], 2, "100")
            if not quantity or line["priceBasis"] not in ("gross", "net"):
                raise ValueError("Количество должно быть положительным; укажите базу цены")
            _tax(line["tax"])
            base = int((quantity * price * (1 - discount / 100) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))
            rate = line["tax"].get("rate", 0)
            vat = int(
                (Decimal(base) * rate / (100 if line["priceBasis"] == "net" else 100 + rate)).quantize(
                    Decimal(1), rounding=ROUND_HALF_UP
                )
            )
            expected = {
                "netMinor": base if line["priceBasis"] == "net" else base - vat,
                "vatMinor": vat,
                "grossMinor": base + vat if line["priceBasis"] == "net" else base,
            }
            key = json.dumps(line["tax"], sort_keys=True)
            group = groups.setdefault(key, {"tax": line["tax"], "netMinor": 0, "vatMinor": 0, "grossMinor": 0})
            for field, amount in expected.items():
                _minor(line[field])
                if int(line[field]) != amount:
                    raise ValueError("Сумма строки не совпадает с proposal-pricing/1")
                totals[field] += amount
                group[field] += amount
    public_totals = _object(document["totals"], "pricingVersion netMinor vatMinor grossMinor byTax")
    if public_totals["pricingVersion"] != "proposal-pricing/1":
        raise ValueError("Неподдерживаемая версия расчёта КП")
    for field, amount in totals.items():
        _minor(public_totals[field])
        if int(public_totals[field]) != amount:
            raise ValueError("Общий итог КП не совпадает с суммой строк")
    expected_groups = [
        {key: str(value) if key != "tax" else value for key, value in group.items()} for group in groups.values()
    ]
    if public_totals["byTax"] != expected_groups:
        raise ValueError("Разбивка НДС не совпадает с суммой строк")
    terms = _object(document["terms"], "delivery payment introduction conclusion")
    for value in terms.values():
        _text(value)
    layout = _object(document["layout"], "style accentColor show footer", "logo")
    if (
        layout["style"] not in ("standard", "compact")
        or not isinstance(layout["accentColor"], str)
        or not re.fullmatch("#[0-9a-fA-F]{6}", layout["accentColor"])
    ):
        raise ValueError("Некорректное оформление КП")
    show = _object(layout["show"], "address requisites contact signer")
    if any(type(value) is not bool for value in show.values()):
        raise ValueError("Некорректные настройки блоков КП")
    _text(layout["footer"])
    if "logo" in layout:
        _asset(layout["logo"])
        if (
            layout["logo"]["mimeType"] not in ("image/png", "image/jpeg")
            or layout["logo"]["sizeBytes"] > 5 * 1024 * 1024
        ):
            raise ValueError("Логотип должен быть PNG/JPEG до 5 МБ")
    if not isinstance(document["attachments"], list) or len(document["attachments"]) > 100:
        raise ValueError("Не более 100 приложений КП")
    for asset in document["attachments"]:
        _asset(asset)
    return document


def _money(minor: str) -> str:
    amount = int(minor)
    return f"{amount // 100:,}".replace(",", " ") + f",{amount % 100:02d}"


def _number(value: str, minimum_decimals=0) -> str:
    whole, _, fraction = value.partition(".")
    fraction = fraction.ljust(minimum_decimals, "0")
    return f"{int(whole):,}".replace(",", " ") + ("," + fraction if fraction else "")


def _tax_label(tax):
    return "Без НДС" if tax["kind"] == "none" else f"НДС {tax['rate']}%"


def _xml(tag, **attrs):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(key), str(value))
    return element


def _paragraph(document, text, *, bold=False, keep=False, style=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = keep
    paragraph.add_run(text).bold = bold
    return paragraph


def _party_text(party, show):
    items = [party["name"]]
    if show["requisites"]:
        items.extend(
            f"{label} {party[key]}" for key, label in (("inn", "ИНН"), ("kpp", "КПП"), ("ogrn", "ОГРН")) if party[key]
        )
    if show["address"] and party["address"]:
        items.append(party["address"])
    if show["contact"] and party["contact"]:
        items.append(party["contact"])
    return " · ".join(items)


def render_docx(document: dict, output: Path, assets: dict[str, Path]) -> None:
    validate_document(document)
    doc = Document()
    compact = document["layout"]["style"] == "compact"
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18 if compact else 20)
    section.left_margin = section.right_margin = Mm(18)
    section.header_distance = section.footer_distance = Mm(9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.space_after = Pt(5 if compact else 7)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 20), ("Heading 1", 13), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor.from_string(document["layout"]["accentColor"][1:])
        style.paragraph_format.keep_with_next = True
        if name == "Title":
            style.font.color.rgb = RGBColor(0, 0, 0)
            for border in style.element.xpath("w:pPr/w:pBdr"):
                border.getparent().remove(border)
    props = doc.core_properties
    props.author = props.last_modified_by = ""
    props.title, props.subject = document["title"], "Коммерческое предложение"
    props.comments = props.keywords = props.category = props.identifier = props.language = ""
    props.revision = document["revision"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    if document["layout"]["footer"]:
        footer.add_run(document["layout"]["footer"] + "\n")
    footer.add_run("Страница ")
    footer._p.append(_xml("w:fldSimple", **{"w:instr": "PAGE"}))
    for run in footer.runs:
        run.font.size = Pt(9)
    logo = document["layout"].get("logo")
    if logo:
        source = assets[logo["sha256"]]
        with Image.open(source) as original:
            if (
                original.format not in ("PNG", "JPEG")
                or original.width * original.height > 16000000
                or max(original.size) > 8192
                or min(original.size) < 1
            ):
                raise ValueError("Некорректные размеры или формат логотипа")
            original.load()
            clean = ImageOps.exif_transpose(original).convert("RGBA")
            clean.info.clear()
            clean.thumbnail((1600, 800))
            clean_path = output.parent / "logo-clean.png"
            clean.save(clean_path, format="PNG")
        logo_width = min(40, 20 * clean.width / clean.height)
        doc.add_picture(str(clean_path), width=Mm(logo_width))
    show = document["layout"]["show"]
    _paragraph(doc, _party_text(document["issuer"], show), bold=True, keep=True)
    _paragraph(doc, "Коммерческое предложение", style="Title")
    _paragraph(
        doc, f"№ {document['number']} от {document['documentDate']} · Редакция {document['revision']}", keep=True
    )
    _paragraph(doc, document["title"], bold=True, keep=True)
    _paragraph(doc, "Заказчик: " + _party_text(document["recipient"], show))
    addressee = document["addressee"]
    if any(addressee.values()):
        _paragraph(doc, "Адресат: " + ", ".join(value for value in addressee.values() if value))
    if document["terms"]["introduction"]:
        _paragraph(doc, document["terms"]["introduction"])
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [12, 69, 21, 35, 37]
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Mm(width)
    table._tbl.tblPr.append(_xml("w:tblLayout", **{"w:type": "fixed"}))
    margins = _xml("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        margins.append(_xml(f"w:{side}", **{"w:w": "75", "w:type": "dxa"}))
    table._tbl.tblPr.append(margins)
    table.rows[0]._tr.get_or_add_trPr().append(_xml("w:tblHeader"))
    for cell, title in zip(table.rows[0].cells, ["№", "Наименование", "Кол-во", "Цена, ₽", "Сумма, ₽"], strict=True):
        cell.text = title
        cell._tc.get_or_add_tcPr().append(_xml("w:shd", **{"w:fill": document["layout"]["accentColor"][1:]}))
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, line in enumerate(document["lines"], 1):
        row = table.add_row()
        tax_info = _tax_label(line["tax"]) + (
            ", цена без НДС" if line["priceBasis"] == "net" and line["tax"]["kind"] == "vat" else ""
        )
        if line["discountPercent"] != "0":
            tax_info += f" · Скидка {_number(line['discountPercent'])}%"
        details = "\n".join(value for value in (line["title"], line["description"], tax_info) if value)
        if len(details) < 800:
            row._tr.get_or_add_trPr().append(_xml("w:cantSplit"))
        values = [
            str(index),
            details,
            _number(line["quantity"]) + "\n" + line["unit"],
            _number(line["unitPrice"], 2),
            _money(line["grossMinor"]),
        ]
        for column_index, (cell, width, value) in enumerate(zip(row.cells, widths, values, strict=True)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index == 1 else WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                run.font.size = Pt(10 if compact else 10.5)
            if index % 2 == 0:
                cell._tc.get_or_add_tcPr().append(_xml("w:shd", **{"w:fill": "F3F5F4"}))
    _paragraph(doc, "", keep=True)
    _paragraph(doc, f"Стоимость без НДС: {_money(document['totals']['netMinor'])} ₽", keep=True)
    for group in document["totals"]["byTax"]:
        suffix = (
            f": {_money(group['vatMinor'])} ₽"
            if group["tax"]["kind"] == "vat"
            else f": {_money(group['grossMinor'])} ₽"
        )
        _paragraph(doc, _tax_label(group["tax"]) + suffix, keep=True)
    _paragraph(doc, f"Итого к оплате: {_money(document['totals']['grossMinor'])} ₽", bold=True)
    for label, key in (("Сроки работ и поставки", "delivery"), ("Условия оплаты", "payment")):
        if document["terms"][key]:
            _paragraph(doc, label, style="Heading 1")
            _paragraph(doc, document["terms"][key])
    _paragraph(doc, "Предложение действительно до " + document["validUntil"])
    if document["terms"]["conclusion"]:
        _paragraph(doc, document["terms"]["conclusion"])
    if show["requisites"] and document["issuer"]["paymentDetails"]:
        _paragraph(doc, "Банковские реквизиты", style="Heading 1")
        _paragraph(doc, document["issuer"]["paymentDetails"])
    if show["contact"] and any(document["contact"].values()):
        _paragraph(doc, "Контакт: " + ", ".join(value for value in document["contact"].values() if value))
    if show["signer"] and document.get("signer", {}).get("fullName"):
        signer = document["signer"]
        _paragraph(
            doc,
            ", ".join(value for value in (signer["position"], signer["fullName"]) if value),
            keep=bool(signer["basis"]),
        )
        if signer["basis"]:
            _paragraph(doc, "Основание полномочий: " + signer["basis"])
    if document["attachments"]:
        _paragraph(doc, "Приложения к предложению", style="Heading 1")
        for index, asset in enumerate(document["attachments"], 1):
            _paragraph(doc, f"{index}. {asset['fileName']}")
    doc.save(output)


def _verified_assets(document, entries, workdir: Path) -> dict[str, Path]:
    references = document["attachments"] + ([document["layout"]["logo"]] if document["layout"].get("logo") else [])
    expected = {asset["sha256"]: asset for asset in references}
    if not isinstance(entries, list):
        raise ValueError("Некорректный список вложений")
    result = {}
    for entry in entries:
        _object(entry, "sha256 path")
        path = Path(entry["path"]).resolve(strict=True)
        if path.parent != (workdir / "assets").resolve() or not path.is_file() or entry["sha256"] not in expected:
            raise ValueError("Вложение не принадлежит текущей задаче КП")
        info = expected[entry["sha256"]]
        if path.stat().st_size != info["sizeBytes"] or _digest(path) != entry["sha256"]:
            raise ValueError("Вложение изменилось после проверки")
        result[entry["sha256"]] = path
    if result.keys() != expected.keys():
        raise ValueError("Не найдено одно из явно выбранных приложений")
    return result


def _digest(path: Path) -> str:
    with path.open("rb") as handle:
        digest = hashlib.sha256()
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
        return digest.hexdigest()


def _render(config: dict, stopped: threading.Event) -> dict:
    document = validate_document(config["document"])
    workdir = Path(config["workdir"]).resolve(strict=True)

    def cancelled():
        return stopped.is_set() or (workdir / "cancel.requested").exists()

    def check_cancelled():
        if cancelled():
            raise CancelledError("Экспорт КП отменён")

    check_cancelled()
    output_format = config["format"]
    if output_format not in ("docx", "pdf", "preview", "zip"):
        raise ValueError("Неподдерживаемый формат КП")
    assets = _verified_assets(document, config.get("assets", []), workdir)
    docx_path = workdir / "proposal.docx"
    render_docx(document, docx_path, assets)
    check_cancelled()
    output = docx_path
    page_count = None
    previews = []
    if output_format != "docx":
        soffice = find_soffice()
        if not soffice:
            raise ValueError("Встроенный офисный модуль недоступен. DOCX можно сохранить отдельно; PDF не создан")
        from scandocument.office_engine import convert_with_office
        from pypdf import PdfReader, PdfWriter

        pdf = workdir / "proposal.pdf"
        office_workdir = workdir / "office"
        office_workdir.mkdir(mode=0o700)
        convert_with_office(docx_path, pdf, soffice, office_workdir, cancelled)
        check_cancelled()
        # Strip automatic Office metadata and reconstruct pages without source paths or authors.
        reader, writer = PdfReader(pdf), PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.add_metadata(
            {"/Title": document["title"], "/Subject": "Коммерческое предложение", "/Creator": "СБК Инструменты"}
        )
        clean = workdir / "proposal-clean.pdf"
        with clean.open("xb") as handle:
            writer.write(handle)
        clean.replace(pdf)
        page_count = len(reader.pages)
        output = pdf
        if output_format == "preview":
            import pypdfium2 as pdfium

            with pdfium.PdfDocument(pdf) as rendered:
                for index in range(len(rendered)):
                    check_cancelled()
                    image_path = workdir / f"page-{index + 1}.png"
                    page = rendered[index]
                    bitmap = page.render(scale=1.2)
                    image = bitmap.to_pil()
                    image.save(image_path)
                    image.close()
                    bitmap.close()
                    page.close()
                    previews.append(str(image_path))
        if output_format == "zip":
            output = workdir / "proposal.zip"
            manifest = {"schemaVersion": 1, "number": document["number"], "revision": document["revision"], "files": []}
            with ZipFile(output, "x", compression=ZIP_DEFLATED) as archive:
                archive.write(docx_path, "Коммерческое предложение.docx")
                archive.write(pdf, "Коммерческое предложение.pdf")
                for index, asset in enumerate(document["attachments"], 1):
                    check_cancelled()
                    member = f"Приложения/{index:03d}-{asset['fileName']}"
                    archive.write(assets[asset["sha256"]], member)
                    manifest["files"].append(
                        {"name": member, "sizeBytes": asset["sizeBytes"], "sha256": asset["sha256"]}
                    )
                archive.writestr("Перечень.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    check_cancelled()
    return {
        "type": "complete",
        "outputPath": str(output),
        "outputBytes": output.stat().st_size,
        "sha256": _digest(output),
        "pageCount": page_count,
        "previewPages": previews,
    }


def render(config: dict) -> dict:
    stopped = threading.Event()
    previous = None
    if threading.current_thread() is threading.main_thread():
        previous = signal.signal(signal.SIGTERM, lambda *_args: stopped.set())
    try:
        return _render(config, stopped)
    finally:
        if previous is not None:
            signal.signal(signal.SIGTERM, previous)
